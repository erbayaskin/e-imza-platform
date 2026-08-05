from __future__ import annotations

from pathlib import Path
from datetime import date
from textwrap import dedent

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "pdf"
TMP = ROOT / "tmp" / "pdfs"
DOCX = OUT / "E_Imza_API_Entegrasyon_Kilavuzu.docx"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F3F7FB"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "6B7280"
INK = "17202A"
WHITE = "FFFFFF"
GREEN = "1F6E43"
GOLD = "8A6500"
RED = "9B1C1C"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[index] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def set_run_font(run, name="Calibri", size=None, bold=None, italic=None, color=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8)
    code.font.color.rgb = RGBColor.from_string("1F2937")
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.12)
    code.paragraph_format.space_before = Pt(4)
    code.paragraph_format.space_after = Pt(8)
    code.paragraph_format.line_spacing = 1.0

    small = doc.styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    small.font.name = "Calibri"
    small.font.size = Pt(9)
    small.font.color.rgb = RGBColor.from_string(MID_GRAY)
    small.paragraph_format.space_after = Pt(4)
    small.paragraph_format.line_spacing = 1.1


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def configure_page(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("E-İmza API | Entegrasyon Kılavuzu")
    set_run_font(r, size=8.5, color=MID_GRAY)
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ErbayProject  |  Ağustos 2026  |  Sayfa ")
    set_run_font(r, size=8.5, color=MID_GRAY)
    add_field(p, "PAGE")


def add_title_page(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("TEKNİK ENTEGRASYON KILAVUZU")
    set_run_font(r, size=11, bold=True, color=GOLD)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("E-İmza API Platformu")
    set_run_font(r, size=30, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run("Server-side ve client-side imzalama, doğrulama, demo ve offline masaüstü uygulaması")
    set_run_font(r, size=14, color=DARK_BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Sürüm 0.1.0-SNAPSHOT")
    set_run_font(r, size=11, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("5 Ağustos 2026")
    set_run_font(r, size=10.5, color=MID_GRAY)
    for _ in range(4):
        doc.add_paragraph()
    callout(doc,
        "Kapsam notu",
        "Bu belge mevcut proje kodu ve OpenAPI sözleşmesine göre hazırlanmıştır. Gerçek ESHS/TSA uçları, üretici sürücüleri, bağımsız birlikte çalışabilirlik ve hukuk/bilgi güvenliği kabulü tamamlanmadan mevzuata tam uygunluk iddiası oluşturmaz.",
        kind="caution")
    doc.add_page_break()


def add_heading(doc, text: str, level=1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc, text: str, bold_prefix: str | None = None, style=None) -> None:
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        r.bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)


def bullets(doc, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def numbered(doc, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def callout(doc, title: str, text: str, kind="info") -> None:
    fill = {"info": PALE_BLUE, "caution": "FFF8E6", "risk": "FDECEC", "success": "EAF6EF"}[kind]
    color = {"info": DARK_BLUE, "caution": GOLD, "risk": RED, "success": GREEN}[kind]
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10.5, bold=True, color=color)
    p = cell.add_paragraph(text)
    p.paragraph_format.space_after = Pt(0)


def code(doc, text: str, language: str | None = None) -> None:
    if language:
        p = doc.add_paragraph(style="Small Note")
        p.paragraph_format.space_before = Pt(4)
        r = p.add_run(language)
        r.bold = True
    p = doc.add_paragraph(style="Code Block")
    p.paragraph_format.keep_together = True
    p.add_run(dedent(text).strip("\n"))
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F4F6F9")
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)


def table(doc, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for idx, value in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_run_font(r, size=9.5, bold=True, color=NAVY)
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=9.25, color=INK)
    set_table_geometry(t, widths)
    doc.add_paragraph(style="Small Note")


def font(size: int, bold=False):
    try:
        return ImageFont.truetype("arialbd.ttf" if bold else "arial.ttf", size)
    except OSError:
        return ImageFont.load_default()


def flow_diagram(path: Path, client: bool) -> None:
    width, height = 1500, 520
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    title = "CLIENT-SIDE AKIŞ" if client else "SERVER-SIDE AKIŞ"
    draw.text((55, 35), title, fill="#17365D", font=font(34, True))
    if client:
        boxes = [
            ("İş Uygulaması", "Oturumu açar\nmanifest alır"),
            ("Merkez API", "Belgeyi hazırlar\nmanifesti imzalar"),
            ("Yerel Agent", "Manifesti doğrular\nPIN'i yerelde ister"),
            ("Akıllı Kart", "Özel anahtarla\nham imza üretir"),
        ]
    else:
        boxes = [
            ("İş Uygulaması", "Oturumu açar\nserver-sign çağırır"),
            ("Merkez API", "serverKeyId profilini\nve politikayı uygular"),
            ("PKCS#11", "Slot/anahtar seçer\nimza üretir"),
            ("Kart veya HSM", "Özel anahtar\ncihazdan çıkmaz"),
        ]
    x_positions = [55, 420, 785, 1150]
    for idx, (name, detail) in enumerate(boxes):
        x = x_positions[idx]
        draw.rounded_rectangle((x, 145, x + 290, 350), radius=18, fill="#F3F7FB", outline="#2E74B5", width=4)
        draw.text((x + 24, 175), name, fill="#17365D", font=font(24, True))
        draw.multiline_text((x + 24, 225), detail, fill="#17202A", font=font(19), spacing=8)
        if idx < len(boxes) - 1:
            start, end = x + 295, x_positions[idx + 1] - 8
            draw.line((start, 248, end, 248), fill="#8A6500", width=6)
            draw.polygon([(end, 248), (end - 22, 234), (end - 22, 262)], fill="#8A6500")
    note = "PIN tarayıcıya ve merkeze gitmez" if client else "HSM PIN'i credentialRef ile secret ortamından alınır"
    draw.text((55, 415), note, fill="#8A6500", font=font(22, True))
    image.save(path)


def add_picture(doc, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.45))
    p = doc.add_paragraph(style="Small Note")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(caption)
    r.italic = True


def add_java_library_markdown(doc, source: Path) -> None:
    """Render the maintained direct-JAR Markdown guide into the handbook."""
    lines = source.read_text(encoding="utf-8").splitlines()
    paragraph: list[str] = []
    code_lines: list[str] = []
    in_code = False
    code_label = "Kod"

    def flush_paragraph() -> None:
        if paragraph:
            add_para(doc, " ".join(part.strip() for part in paragraph))
            paragraph.clear()

    for raw in lines:
        line = raw.rstrip()
        if line.startswith(chr(96) * 3):
            if in_code:
                code(doc, "\n".join(code_lines), code_label)
                code_lines.clear()
                in_code = False
            else:
                flush_paragraph()
                code_label = line[3:].strip() or "Kod"
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line.startswith("# "):
            flush_paragraph()
            continue
        if line.startswith("## "):
            flush_paragraph()
            add_heading(doc, line[3:], 2)
            continue
        if line.startswith("- "):
            flush_paragraph()
            doc.add_paragraph(line[2:], style="List Bullet")
            continue
        if len(line) > 3 and line[0].isdigit() and ". " in line[:4]:
            flush_paragraph()
            doc.add_paragraph(line.split(". ", 1)[1], style="List Number")
            continue
        if not line.strip():
            flush_paragraph()
            continue
        paragraph.append(line)
    flush_paragraph()

def build() -> Path:
    OUT.mkdir(parents=True, exist_ok=True)
    TMP.mkdir(parents=True, exist_ok=True)
    client_img = TMP / "client-side-flow.png"
    server_img = TMP / "server-side-flow.png"
    flow_diagram(client_img, True)
    flow_diagram(server_img, False)

    doc = Document()
    configure_styles(doc)
    configure_page(doc)
    add_title_page(doc)

    add_heading(doc, "İçindekiler", 1)
    for item in [
        "1. Kılavuzun amacı ve hedef kitle",
        "2. Platformu tanıyın",
        "3. İmza formatları ve seviyeler",
        "4. Hızlı başlangıç ve demo uygulamaları",
        "5. Ortak API kuralları",
        "6. Server-side entegrasyon",
        "7. Client-side entegrasyon",
        "8. İmza ve sertifika doğrulama",
        "9. Güven deposu, TSA ve politikalar",
        "10. Hata yönetimi ve güvenlik",
        "11. Canlı ortam kontrol listesi",
        "12. Doğrudan Java/JAR entegrasyonu",
        "13. Çoklu imza: REST ve demo",
        "Ek A. Uç noktalar",
        "Ek B. Yardımcı kod parçaları",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    callout(doc, "Okuma yolu", "REST entegrasyonu için 4 ve 5. bölümlerden sonra server-side için 6, client-side için 7. bölüme geçin. REST kullanmadan JAR'ları doğrudan çağıracaksanız 12; seri/paralel imza için 13. bölümü okuyun.")
    doc.add_page_break()

    add_heading(doc, "1. Kılavuzun amacı ve hedef kitle", 1)
    add_para(doc, "Bu kılavuz E-İmza API platformunu çalıştıracak ekipler, kurumsal uygulamalarına imzalama ekleyecek geliştiriciler, akıllı kart/HSM yöneticileri ve doğrulama hizmetini kullanacak servis ekipleri için hazırlanmıştır.")
    bullets(doc, [
        "Platformun modüllerini ve güvenlik sınırlarını anlamak",
        "Web demo, yönetim ekranı, Smart Card Agent ve offline masaüstünü kullanmak",
        "Server-side kart veya HSM ile imza üretmek",
        "Client-side Smart Card Agent ile son kullanıcı kartından imza almak",
        "CAdES, XAdES ve PAdES imzalarını ve imzalayan sertifikayı doğrulamak",
        "Güvenilir kök/alt kök, RFC 3161 TSA ve doğrulama politikalarını yönetmek",
    ])
    callout(doc, "Önemli", "Özel anahtar hiçbir akışta karttan veya HSM'den dışarı çıkarılmaz. Sertifika public veridir; indirilen .cer dosyası özel anahtar veya PIN içermez.", "success")

    add_heading(doc, "2. Platformu tanıyın", 1)
    add_para(doc, "Platform Java 21 ve Spring Boot tabanlı, çok modüllü bir elektronik imza altyapısıdır. İmzalama çekirdeği ile cihaz erişimi ayrıdır; böylece aynı format üretim motoru server-side ve client-side akışlarda kullanılabilir.")
    table(doc, ["Modül", "Görev"], [
        ["signature-core", "Ortak modeller, politika ve imzalama oturumu durum makinesi"],
        ["certificate-validation", "PKIX sertifika yolu, tarih, NES/politika, CRL ve OCSP kontrolleri"],
        ["timestamp-client", "RFC 3161 zaman damgası isteği ve token doğrulaması"],
        ["signature-cades", "CAdES B-B/B-T, paralel/seri imza, B-LT/B-LTA ve doğrulama"],
        ["signature-xades", "Detached/enveloped/enveloping, paralel/seri XAdES ve doğrulama"],
        ["signature-pades", "PDF ByteRange, ETSI.CAdES.detached ve incremental seri PAdES"],
        ["smartcard-agent", "Loopback PC/SC, ATR, PKCS#11, yerel PIN ve client-side kart imzası"],
        ["signature-api", "REST API, oturum orkestrasyonu, DB, yönetim, demo ve doğrulama"],
        ["desktop-signing-demo", "API/agent gerektirmeyen offline Swing CAdES uygulaması"],
    ], [2400, 6960])

    add_heading(doc, "2.1 Uygulama yüzeyleri", 2)
    table(doc, ["Adres / uygulama", "Amaç"], [
        ["/demo/", "Server-side veya client-side CAdES/XAdES/PAdES imzalama demosu"],
        ["/multi-signature/", "Mevcut imzalı dosyaya paralel veya seri imza demosu"],
        ["/validation/", "İmza doğrulama, imzalayan bilgileri ve .cer indirme"],
        ["/admin/", "Cihaz profili, güven deposu, TSA ve politika yönetimi"],
        ["127.0.0.1:18443/agent/v1", "Yalnız client-side akışta tarayıcı ile yerel agent arasındaki loopback API"],
        ["desktop-signing-demo-...-exec.jar", "Tamamen offline CAdES imzalama ve doğrulama"],
    ], [2850, 6510])

    add_heading(doc, "3. İmza formatları ve seviyeler", 1)
    table(doc, ["Format", "Paketleme", "Çıktı", "Tipik kullanım"], [
        ["CAdES", "DETACHED / ATTACHED", ".p7s", "Her tür dosyanın CMS tabanlı imzalanması"],
        ["XAdES", "DETACHED / ENVELOPED / ENVELOPING", ".xades.xml", "XML iş akışları ve XML içinde imza"],
        ["PAdES", "ENVELOPED", ".signed.pdf", "PDF'nin kendi içinde imza"],
    ], [1300, 2200, 1600, 4260])
    table(doc, ["Seviye", "Anlam", "Gereksinim"], [
        ["B_B", "Temel imza", "İmzalayan sertifikası ve kriptografik imza"],
        ["B_T", "İmza zaman damgalı", "Etkin RFC 3161 TSA ve güvenilir TSA zinciri"],
        ["B_LT", "Uzun dönem kanıtları gömülü", "Tam zincir ve CRL/OCSP kanıtları; CAdES yükseltme"],
        ["B_LTA", "Arşiv zaman damgalı", "B-LT materyali ve arşiv TSA; CAdES yükseltme/yenileme"],
    ], [1100, 3300, 4960])
    callout(doc, "İlk test önerisi", "TSA yapılandırılmadan B_B ile başlayın. B_T seçimi için TSA endpoint'i ve TSA güven kökü hazır olmalıdır.", "caution")

    add_heading(doc, "4. Hızlı başlangıç ve demo uygulamaları", 1)
    add_heading(doc, "4.1 Gereksinimler ve paketleme", 2)
    bullets(doc, ["Java 21", "Maven 3.8.8 veya üzeri", "Local test için H2; üretim için PostgreSQL", "Kart senaryoları için üreticinin onaylı PKCS#11 middleware'i"])
    code(doc, r'''
      cd D:\ErbayProject
      mvn "-Dmaven.repo.local=D:\ErbayProject\.m2\repository" -pl signature-api,smartcard-agent,desktop-signing-demo -am -DskipTests package
    ''', "PowerShell / CMD uyarlaması")
    code(doc, r'''
      java -jar .\signature-api\target\signature-api-0.1.0-SNAPSHOT-exec.jar --spring.profiles.active=local
    ''', "Merkez API")
    add_para(doc, "Sağlık kontrolü: http://localhost:8080/actuator/health")

    add_heading(doc, "4.2 Yönetim ekranı", 2)
    add_para(doc, "http://localhost:8080/admin/ adresi server key profillerini, güvenilir sertifikaları, TSA profilini ve doğrulama politikasını yönetir. Local örnek tenant UUID'si 11111111-1111-1111-1111-111111111111 değeridir.")
    numbered(doc, [
        "Akıllı kart için PC/SC taramasıyla ATR'yi tespit edin veya manuel profil girin.",
        "SMART_CARD için ATR, maske ve mutlak PKCS#11 yolu tanımlayın; slot otomatik bulunur.",
        "HSM için ATR girmeyin; kütüphane yolu, slot ve credentialRef tanımlayın.",
        "Kart sertifikasının kök ve alt kök CA sertifikalarını güven deposuna ekleyin.",
        "B-T ve üzeri için TSA profilini ve TSA güven zincirini tanımlayın.",
        "Ortam gereksinimine göre STRICT, CUSTOM veya AUDIT_ONLY politika modunu seçin.",
    ])

    add_heading(doc, "4.3 Web imzalama ve doğrulama demoları", 2)
    bullets(doc, [
        "İmzalama: http://localhost:8080/demo/",
        "Seri/paralel çoklu imza: http://localhost:8080/multi-signature/",
        "Doğrulama: http://localhost:8080/validation/",
        "Doğrulama sonucu imzalayanın subject/issuer, seri, geçerlilik, algoritma ve parmak izi bilgilerini gösterir.",
        "İmzalayanın public X.509 sertifikası DER .cer olarak indirilebilir.",
    ])

    add_heading(doc, "4.4 Offline masaüstü", 2)
    code(doc, r'''
      java -jar .\desktop-signing-demo\target\desktop-signing-demo-0.1.0-SNAPSHOT-exec.jar
    ''', "PowerShell")
    add_para(doc, "Masaüstü uygulaması merkez API'ye bağlanmaz, HTTP sunucusu ve Smart Card Agent başlatmaz. Kart adı, ATR ve PKCS#11 kütüphane yolu yerel profilde saklanır; PIN saklanmaz. Uygulama attached/detached CAdES imzalar, doğrular ve imzalayan sertifikasını .cer olarak kaydeder.")

    add_heading(doc, "5. Ortak API kuralları", 1)
    add_heading(doc, "5.1 Temel adres ve başlıklar", 2)
    table(doc, ["Öğe", "Kural"], [
        ["Base URL", "https://sunucu.example/api/v1 (local: http://localhost:8080/api/v1)"],
        ["Authorization", "Üretimde OAuth2 Bearer JWT; scope: eimza.sign, eimza.validate veya eimza.admin"],
        ["X-Tenant-Id", "Zorunlu UUID; üretimde doğrulanmış token tenant claim'i ile eşleşmelidir"],
        ["Idempotency-Key", "Oturum oluştururken zorunlu; aynı iş isteğinin tekrarını güvenli yönetir"],
        ["X-Correlation-Id", "İsteğe bağlı; uçtan uca log/olay korelasyonu için önerilir"],
        ["Content-Type", "JSON isteklerinde application/json"],
    ], [2200, 7160])
    add_heading(doc, "5.2 Base64 ve Base64URL", 2)
    bullets(doc, [
        "Belge, sertifika ve artifact alanları standart Base64 kullanır.",
        "documentDigest.value, ham kart signature ve deviceSignature padding'siz Base64URL kullanır.",
        "Karakter dönüşümü yapılmamalı; özet doğrudan dosyanın baytlarından hesaplanmalıdır.",
    ])
    code(doc, r'''
      static String b64(byte[] value) {
          return Base64.getEncoder().encodeToString(value);
      }
      static String b64url(byte[] value) {
          return Base64.getUrlEncoder().withoutPadding().encodeToString(value);
      }
    ''', "Java 21")
    add_heading(doc, "5.3 Oturum yaşam döngüsü", 2)
    table(doc, ["Aşama", "Client-side", "Server-side"], [
        ["1", "POST /signing-sessions", "POST /signing-sessions"],
        ["2", "POST /{id}/manifest", "POST /{id}/server-sign"],
        ["3", "agent-connected ve approve", "API cihaz profilini çözer"],
        ["4", "Agent /signing-requests", "PKCS#11 kart/HSM imzası"],
        ["5", "POST /{id}/complete", "Artifact aynı çağrıda döner"],
        ["6", "GET /{id}/artifact", "GET /{id}/artifact"],
    ], [900, 4230, 4230])

    add_heading(doc, "6. Server-side entegrasyon", 1)
    add_picture(doc, server_img, "Şekil 1 - Server-side imzalama güven sınırı")
    add_para(doc, "Server-side modda kart veya HSM, signature-api ile aynı güvenli sunucu ortamında erişilebilir durumdadır. İş uygulaması sürücü yolu veya slot göndermez; yalnız yönetici tarafından önceden oluşturulmuş serverKeyId kullanır.")
    add_heading(doc, "6.1 Ne zaman seçilmeli?", 2)
    bullets(doc, [
        "Kurumsal mühür veya servis imzası merkezi bir HSM'de tutuluyorsa",
        "İmza işlemi kullanıcının bilgisayarından bağımsız yürümeliyse",
        "Sunucuya fiziksel akıllı kart ve okuyucu bağlanmışsa",
        "Yüksek hacim için HSM partition/slot yönetimi gerekiyorsa",
    ])
    add_heading(doc, "6.2 Ön hazırlık", 2)
    table(doc, ["Cihaz", "Profil kuralları", "PIN yönetimi"], [
        ["SMART_CARD", "ATR + maske + mutlak PKCS#11 yolu; autoDiscoverSlot=true", "Tek imza isteğinde pin alanı; DB/log'a yazılmaz"],
        ["HSM", "Mutlak PKCS#11 yolu + sabit slot; ATR yok", "HTTP'de PIN yok; credentialRef ile ortam secret'ından"],
    ], [1700, 4700, 2960])

    add_heading(doc, "6.3 Adım 1 - belge özeti ve oturum", 2)
    code(doc, r'''
      byte[] document = Files.readAllBytes(Path.of("sozlesme.pdf"));
      byte[] digest = MessageDigest.getInstance("SHA-256").digest(document);

      String requestJson = """
      {
        "documentId":"%s",
        "documentDigest":{"algorithm":"SHA-256","value":"%s"},
        "documentName":"sozlesme.pdf",
        "mediaType":"application/pdf",
        "size":%d,
        "format":"PADES",
        "targetLevel":"B_B",
        "turkishProfile":"P1",
        "purpose":"Sözleşme onayı",
        "signingMode":"SERVER_SIDE",
        "deviceId":null,
        "serverKeyId":"%s",
        "documentBase64":"%s",
        "signaturePackaging":"ENVELOPED",
        "signatureAlgorithm":"RSA_PKCS1_SHA256"
      }
      """.formatted(UUID.randomUUID(), b64url(digest), document.length,
                    serverKeyId, b64(document));
    ''', "Java - istek gövdesi")
    code(doc, r'''
      HttpRequest create = HttpRequest.newBuilder(base.resolve("/api/v1/signing-sessions"))
          .header("Authorization", "Bearer " + accessToken)
          .header("X-Tenant-Id", tenantId)
          .header("Idempotency-Key", UUID.randomUUID().toString())
          .header("X-Correlation-Id", UUID.randomUUID().toString())
          .header("Content-Type", "application/json")
          .POST(HttpRequest.BodyPublishers.ofString(requestJson))
          .build();
      HttpResponse<String> created = client.send(create,
          HttpResponse.BodyHandlers.ofString());
      // 201 cevabındaki sessionId JSON kütüphanesiyle okunur.
    ''', "Java 21 HttpClient")

    add_heading(doc, "6.4 Adım 2 - sunucuda imzala", 2)
    code(doc, r'''
      POST /api/v1/signing-sessions/{sessionId}/server-sign
      Authorization: Bearer <token>
      X-Tenant-Id: 11111111-1111-1111-1111-111111111111
      Content-Type: application/json

      { "pin": "yalnız-smart-card-icin-tek-kullanimlik" }
    ''', "HTTP - SMART_CARD")
    code(doc, r'''
      POST /api/v1/signing-sessions/{sessionId}/server-sign
      ...
      { "pin": null }
    ''', "HTTP - HSM")
    add_para(doc, "Başarılı cevap SignatureArtifact döndürür: artifactId, format, level, mediaType, sha256 ve artifactBase64. Base64 çözülerek .p7s, .xades.xml veya .signed.pdf yazılır.")
    code(doc, r'''
      byte[] artifact = Base64.getDecoder().decode(artifactBase64);
      Files.write(Path.of("sozlesme.signed.pdf"), artifact);
    ''', "Java")
    callout(doc, "Server-side güvenlik", "İş uygulaması PKCS#11 DLL yolu, HSM slotu veya credentialRef seçemez. Bu değerler eimza.admin yetkili yönetici tarafından önceden tanımlanır.", "risk")

    doc.add_page_break()
    add_heading(doc, "7. Client-side entegrasyon", 1)
    add_picture(doc, client_img, "Şekil 2 - Client-side Smart Card Agent akışı")
    add_para(doc, "Client-side mod yalnız akıllı kart içindir; HSM desteklenmez. Tarayıcı WebCrypto ile belge özetini hesaplar, merkezden imzalı manifest alır ve loopback agent'a iletir. PIN yalnız agent'ın Swing penceresinde girilir.")
    add_heading(doc, "7.1 Agent kurulumu ve yapılandırması", 2)
    code(doc, r'''
      $env:EIMZA_AGENT_DEVICE_ID = "a8a0dc09-54ab-40b7-b404-bebd55ff1756"
      $env:EIMZA_AGENT_MANIFEST_PUBLIC_KEY = (
        Invoke-RestMethod http://localhost:8080/api/v1/signing-configuration/manifest-key
      ).publicKey
      $env:EIMZA_AGENT_ALLOWED_ORIGIN = "https://uygulama.example"
      $env:SPRING_CONFIG_ADDITIONAL_LOCATION = "file:D:/eimza/agent.yml"

      java -jar smartcard-agent-0.1.0-SNAPSHOT-exec.jar
    ''', "PowerShell")
    callout(doc, "Origin sınırı", "EIMZA_AGENT_ALLOWED_ORIGIN yalnız gerçek iş uygulaması origin'ini içermelidir. Joker (*) veya rastgele web sitelerine izin verilmemelidir.", "risk")

    add_heading(doc, "7.2 Cihaz kaydı", 2)
    code(doc, r'''
      const agent = 'http://127.0.0.1:18443/agent/v1';
      const identity = await fetch(agent + '/device').then(r => r.json());

      await fetch('/api/v1/admin/client-devices', {
        method: 'POST',
        headers: {
          'Authorization': `Bearer ${token}`,
          'X-Tenant-Id': tenantId,
          'Content-Type': 'application/json'
        },
        body: JSON.stringify({
          deviceId: identity.deviceId,
          displayName: 'Muhasebe bilgisayarı',
          publicKey: identity.publicKey
        })
      });
    ''', "JavaScript")
    add_para(doc, "Üretimde cihaz kaydı normal son kullanıcı yetkisinden ayrılmalı; yönetici onayı, cihaz envanteri ve iptal süreci uygulanmalıdır. Ed25519 özel cihaz anahtarı agent'tan çıkmaz.")

    add_heading(doc, "7.3 Kart ve public sertifika seçimi", 2)
    code(doc, r'''
      const cards = await fetch(agent + '/cards').then(r => r.json());
      const card = cards.find(item => item.status === 'READY');
      if (!card) throw new Error('Hazır kart bulunamadı');

      const certificates = await fetch(
        agent + '/cards/' + encodeURIComponent(card.readerId) + '/certificates/refresh',
        { method: 'POST', headers: { 'X-EImza-Agent': '1' } }
      ).then(r => r.json());
      const signer = certificates.find(item => item.hasPrivateKey);
    ''', "JavaScript")
    add_para(doc, "Public sertifikalar mümkünse PIN'siz listelenir. Middleware public nesnelere oturumsuz erişim vermiyorsa üretici adaptörünün davranışına göre yerel PIN fallback'i gerekebilir. PIN hiçbir HTTP gövdesinde kabul edilmez.")

    add_heading(doc, "7.4 Oturum, manifest ve kart imzası", 2)
    code(doc, r'''
      const bytes = new Uint8Array(await file.arrayBuffer());
      const digest = new Uint8Array(await crypto.subtle.digest('SHA-256', bytes));
      const session = await apiJson('/signing-sessions', {
        method: 'POST',
        headers: { 'Idempotency-Key': crypto.randomUUID() },
        body: JSON.stringify({
          documentId: crypto.randomUUID(),
          documentDigest: { algorithm: 'SHA-256', value: b64url(digest) },
          documentName: file.name, mediaType: file.type, size: file.size,
          format: 'CADES', targetLevel: 'B_B', turkishProfile: 'P1',
          purpose: 'Kullanıcı onaylı imza', signingMode: 'CLIENT_SIDE',
          deviceId: identity.deviceId, serverKeyId: null,
          documentBase64: b64(bytes), signaturePackaging: 'DETACHED',
          signatureAlgorithm: 'RSA_PKCS1_SHA256'
        })
      });
    ''', "JavaScript - oturum")
    code(doc, r'''
      const manifest = await apiJson(`/signing-sessions/${session.sessionId}/manifest`, {
        method: 'POST',
        body: JSON.stringify({
          readerId: card.readerId,
          certificateFingerprint: signer.fingerprintSha256,
          certificateBase64: signer.certificateBase64,
          signatureAlgorithm: 'RSA_PKCS1_SHA256'
        })
      });

      await apiJson(`/signing-sessions/${session.sessionId}/agent-connected`, {method:'POST'});
      await apiJson(`/signing-sessions/${session.sessionId}/approve`, {method:'POST'});

      const signed = await fetch(agent + '/signing-requests', {
        method: 'POST',
        headers: {'Content-Type':'application/json','X-EImza-Agent':'1'},
        body: JSON.stringify({manifest: manifest.manifest, signature: manifest.signature})
      }).then(r => r.json());

      const artifact = await apiJson(`/signing-sessions/${session.sessionId}/complete`, {
        method: 'POST',
        body: JSON.stringify({
          signature: signed.signature,
          deviceSignature: signed.deviceSignature
        })
      });
    ''', "JavaScript - manifest ve tamamlama")
    callout(doc, "Kullanıcı onayı", "Belge adı, boyutu, özet/amaç ve sertifika kullanıcıya PIN penceresinden önce gösterilmelidir. Beklenmeyen içerikte kullanıcı işlemi iptal etmelidir.", "caution")

    add_heading(doc, "7.5 apiJson yardımcı fonksiyonu", 2)
    code(doc, r'''
      async function apiJson(path, options = {}) {
        const headers = {
          'Authorization': `Bearer ${token}`,
          'X-Tenant-Id': tenantId,
          'X-Correlation-Id': crypto.randomUUID(),
          'Content-Type': 'application/json',
          ...(options.headers || {})
        };
        const response = await fetch('/api/v1' + path, {...options, headers});
        const body = await response.json();
        if (!response.ok) throw new Error(`${body.code}: ${body.detail || body.message}`);
        return body;
      }
    ''', "JavaScript")

    add_heading(doc, "8. İmza ve sertifika doğrulama", 1)
    add_heading(doc, "8.1 İmza doğrulama isteği", 2)
    code(doc, r'''
      POST /api/v1/validations/signatures
      Authorization: Bearer <token>
      X-Tenant-Id: 11111111-1111-1111-1111-111111111111
      Content-Type: application/json

      {
        "format":"CADES",
        "signaturePackaging":"DETACHED",
        "content":"<orijinal-belge-Base64>",
        "signature":"<p7s-Base64>",
        "validationTime":null
      }
    ''', "HTTP")
    table(doc, ["Format", "signature", "content"], [
        ["CAdES ATTACHED", "DER CMS/CAdES Base64", "null"],
        ["CAdES DETACHED", "DER CMS/CAdES Base64", "Orijinal belge Base64"],
        ["XAdES ENVELOPED/ENVELOPING", "İmzalı XML Base64", "null"],
        ["XAdES DETACHED", "İmzalı XML Base64", "Orijinal belge Base64"],
        ["PAdES", "İmzalı PDF Base64", "null"],
    ], [2200, 3600, 3560])
    add_heading(doc, "8.2 Sonucu yorumlama", 2)
    table(doc, ["Ana sonuç", "Yorum"], [
        ["VALID", "Kriptografik, sertifika ve etkin zorunlu politikalar geçti"],
        ["INVALID", "Kriptografik/biçimsel hata veya kesin politika ihlali var; imza kabul edilmez"],
        ["INDETERMINATE", "Kanıt/hizmet eksik, politika pasif veya güven kararı tamamlanamadı; otomatik geçerli sayılmaz"],
    ], [1800, 7560])
    add_para(doc, "checks dizisi her kontrolün code, indication, message, evidenceDigest ve details alanlarını taşır. Karar verirken yalnız summary metnine değil mainIndication ve checks alanlarına bakın.")

    add_heading(doc, "8.3 İmzalayan bilgileri ve sertifika indirme", 2)
    code(doc, r'''
      const signer = report.signer;
      console.log(signer.commonName, signer.subjectDn, signer.issuerDn);
      console.log(signer.validFrom, signer.validUntil, signer.sha256Fingerprint);

      const der = Uint8Array.from(atob(signer.certificateBase64), c => c.charCodeAt(0));
      const url = URL.createObjectURL(new Blob([der], {type:'application/pkix-cert'}));
      const link = document.createElement('a');
      link.href = url;
      link.download = `imzalayan-${signer.certificateSerialNumber}.cer`;
      link.click();
      URL.revokeObjectURL(url);
    ''', "JavaScript")
    callout(doc, "Sertifika public veridir", "certificateBase64 yalnız DER X.509 public sertifikayı içerir. Bununla imza üretilemez; özel anahtar donanımda kalır.")

    add_heading(doc, "8.4 Sertifika doğrulama", 2)
    code(doc, r'''
      POST /api/v1/validations/certificates
      {
        "certificate":"<leaf-DER-Base64>",
        "intermediateCertificates":["<alt-kok-DER-Base64>"],
        "validationTime":null
      }
    ''', "HTTP")

    add_heading(doc, "9. Güven deposu, TSA ve politikalar", 1)
    add_heading(doc, "9.1 Güvenilir kök ve alt kökler", 2)
    bullets(doc, [
        "Yalnız CA sertifikaları ROOT veya INTERMEDIATE olarak eklenir; son kullanıcı sertifikası trust anchor yapılmaz.",
        "Ekleme/güncelleme/silme yeni bir tarihsel snapshot sürümü üretir; geçmiş kararlar izlenebilir kalır.",
        "Sertifika zinciri imza zamanı veya doğrulama zamanı için ilgili snapshot'a göre kurulmalıdır.",
    ])
    code(doc, r'''
      POST /api/v1/admin/trusted-certificates
      {
        "certificate":"<PEM-veya-DER-Base64>",
        "trustType":"ROOT",
        "displayName":"Kurumsal güvenilir kök",
        "enabled":true
      }
    ''', "HTTP")
    add_heading(doc, "9.2 RFC 3161 TSA", 2)
    bullets(doc, [
        "TSA endpoint, timeout, sağlayıcı ID ve credentialRef yönetim ekranından sürümlenir.",
        "Gerçek Authorization değeri DB'ye değil credentialRef'in işaret ettiği secret ortamına yazılır.",
        "TSA sertifika zincirinin kökü de güven deposunda bulunmalıdır.",
        "B-T/B-LTA üretiminde timestamp policy OID ve token doğrulaması izlenmelidir.",
    ])
    add_heading(doc, "9.3 Doğrulama politikaları", 2)
    table(doc, ["Politika", "Amaç"], [
        ["QC compliance", "Nitelikli sertifika/QC ifadelerini denetler"],
        ["Certificate policy", "Zorunlu sertifika politika OID'lerini denetler"],
        ["Revocation", "CRL/OCSP kanıtını ve erişilebilirliğini denetler"],
        ["Signature policy", "İmza politikası OID/özetini denetler"],
        ["Signing certificate validity", "Sertifikanın imzalama zamanındaki geçerliliğini denetler"],
    ], [3000, 6360])
    callout(doc, "Production kuralı", "Sertifika tarih kontrolü production profilinde zorunlu aktiftir. Local/test ortamında süresi dolmuş kartla yalnız donanım entegrasyon testi için pasifleştirilirse sonuç hukuken geçerli imza sayılmaz.", "risk")

    add_heading(doc, "10. Hata yönetimi ve güvenlik", 1)
    add_heading(doc, "10.1 Problem Details", 2)
    code(doc, r'''
      {
        "type":"https://errors.eimza.local/...",
        "title":"SERVER_PKCS11_SIGNING_FAILED",
        "status":422,
        "detail":"Sunucu PKCS#11 imzalama işlemi tamamlanamadı.",
        "code":"SERVER_PKCS11_SIGNING_FAILED",
        "correlationId":"...",
        "retryable":false,
        "timestamp":"..."
      }
    ''', "Örnek hata")
    table(doc, ["Durum", "Uygulama davranışı"], [
        ["400", "İstek şemasını ve Base64/Base64URL kullanımını düzeltin; otomatik tekrar etmeyin"],
        ["401/403", "Token, scope ve tenant claim eşleşmesini kontrol edin"],
        ["404", "Tenant kapsamındaki session/serverKey kaydını kontrol edin"],
        ["409", "Oturum durumunu GET ile okuyun; idempotency sonucunu kullanın"],
        ["422", "Kart/sertifika/politika/format ayrıntısını code ve checks üzerinden çözün"],
        ["5xx / retryable=true", "Sınırlı exponential backoff ve aynı Idempotency-Key ile tekrar değerlendirin"],
    ], [1500, 7860])
    add_heading(doc, "10.2 Zorunlu güvenlik kontrolleri", 2)
    bullets(doc, [
        "TLS olmadan üretim trafiği taşımayın; loopback agent dışında HTTP kullanmayın.",
        "PIN, HSM secret, token ve özel anahtar materyalini loglamayın.",
        "Client-side PIN'i web formuna koymayın; yalnız agent Swing penceresi kullanmalıdır.",
        "Belge içeriğini ve Base64 artifact'i uygulama loglarına yazmayın.",
        "OAuth scope ve tenant izolasyonunu her uçta uygulayın.",
        "Idempotency-Key'i iş işlemi bazında üretin; rastgele tekrarlar yeni imza oluşturmasın.",
        "Manifest Ed25519 anahtarını üretimde KMS/HSM/secret manager ile sürümleyin.",
        "Agent origin allowlist, imzalı dağıtım ve cihaz iptal prosedürü kullanın.",
        "PKCS#11 sürücülerini yalnız üretici/kurum onaylı kaynaktan kurun.",
    ])

    add_heading(doc, "11. Canlı ortam kontrol listesi", 1)
    numbered(doc, [
        "PostgreSQL, yedekleme ve Flyway migration yetkilerini doğrulayın.",
        "OIDC issuer, audience, JWS algoritmaları ve tenant claim eşleşmesini doğrulayın.",
        "Manifest Ed25519 anahtar çiftini kalıcı ve sürümlü secret olarak sağlayın.",
        "Güvenilir kök/alt kök ve TSA zincirini kurum politikasıyla onaylatın.",
        "OCSP/CRL/TSA ağ çıkışlarını allowlist ve timeout kurallarıyla açın.",
        "HSM slot/partition ve credentialRef değerlerini iki kişi kontrolüyle tanımlayın.",
        "Agent paketini kod imzalı dağıtın; origin ve cihaz kayıt yaşam döngüsünü yönetin.",
        "CAdES/XAdES/PAdES için pozitif, bozuk içerik, yanlış paketleme ve süresi dolmuş sertifika testlerini çalıştırın.",
        "Yük, felaket kurtarma, bağımsız sızma ve fiziksel kart kabul testlerini tamamlayın.",
        "Correlation ID, metrik, alarm ve olay müdahale runbook'larını operasyon ekibine teslim edin.",
    ])

    doc.add_page_break()
    add_heading(doc, "12. Doğrudan Java/JAR entegrasyonu", 1)
    add_para(doc, "Bu bölüm REST uçlarını veya demo ekranlarını kullanmadan signature-* JAR modüllerini kendi Java 21 uygulamanızda çağırmayı gösterir. Örnekler proje ile birlikte sürümlenen JAVA_KUTUPHANE_ENTEGRASYONU.md kaynağından üretilir.")
    callout(doc, "Temel sözleşme", "prepare ile imzalanacak baytları üretin; kart/HSM/PrivateKey ile imzalayın; completeBaseline veya completeWithTimestamp ile standart artifact'i tamamlayın. Özel anahtar format modüllerine verilmez.", "success")
    add_java_library_markdown(doc, ROOT / "JAVA_KUTUPHANE_ENTEGRASYONU.md")

    add_heading(doc, "13. Çoklu imza: REST ve demo", 1)
    table(doc, ["Format", "PARALLEL", "SERIAL", "Paketleme notu"], [
        ["CAdES", "Yeni üst seviye SignerInfo", "Hedef imzada counterSignature", "ATTACHED / DETACHED"],
        ["XAdES", "Bağımsız ds:Signature", "xades:CounterSignature", "Paralel: DETACHED / ENVELOPING"],
        ["PAdES", "Desteklenmez", "Incremental PDF revision", "ENVELOPED"],
    ], [1300, 2700, 2700, 2660])
    code(doc, r"""
      {
        "...":"normal oturum alanları",
        "format":"CADES",
        "multiSignatureType":"PARALLEL",
        "existingArtifactBase64":"<önceki-p7s-Base64>",
        "targetSignatureIndex":0,
        "documentBase64":"<orijinal-belge-Base64>"
      }
    """, "POST /api/v1/signing-sessions")
    bullets(doc, [
        "SINGLE varsayılandır; eski istemciler yeni alanları göndermeden çalışır.",
        "PARALLEL ve SERIAL için existingArtifactBase64 zorunludur.",
        "CAdES/XAdES PARALLEL için orijinal belge gerekir; PAdES yalnız SERIAL kabul eder.",
        "XAdES ENVELOPED + PARALLEL reddedilir; DETACHED veya ENVELOPING seçilir.",
        "targetSignatureIndex sıfır tabanlıdır ve CAdES/XAdES seri imzada hedef imzayı seçer.",
        "İki Base64 alanı için varsayılan gövde sınırı 70 MiB'dir; EIMZA_MAXIMUM_REQUEST_BYTES ile ortam kapasitesine göre değiştirilebilir.",
    ])
    add_heading(doc, "13.1 Demo adımları", 2)
    numbered(doc, [
        "İlk imzayı /demo/ sayfasında üretip artifact'i indirin.",
        "/multi-signature/ sayfasında mevcut imzalı artifact'i seçin.",
        "CAdES/XAdES PARALLEL için orijinal dosyayı da seçin; SERIAL için hedef indeksi girin.",
        "Client-side kart/sertifika veya server-side serverKeyId profilini seçerek imzalayın.",
        "Yeni artifact'i indirin ve /validation/ sayfasında doğrulayın.",
    ])
    callout(doc, "Doğrulama kapsamı", "CAdES'teki tüm üst seviye ve iç içe SignerInfo değerleri, XAdES'teki tüm ds:Signature/CounterSignature öğeleri ve PAdES'teki tüm imza sözlükleri doğrulanır.", "success")

    add_heading(doc, "Ek A. Uç noktalar", 1)
    table(doc, ["Yöntem ve yol", "Amaç"], [
        ["POST /signing-sessions", "İmzalama oturumu oluşturur"],
        ["GET /signing-sessions/{id}", "Oturum durumunu getirir"],
        ["POST /signing-sessions/{id}/manifest", "Client-side imzalı manifest üretir"],
        ["POST /signing-sessions/{id}/agent-connected", "Agent bağlantısını kaydeder"],
        ["POST /signing-sessions/{id}/approve", "Kullanıcı onayı sonrası kart imzasına geçer"],
        ["POST /signing-sessions/{id}/complete", "Client-side ham imzayı tamamlar"],
        ["POST /signing-sessions/{id}/server-sign", "Server-side kart/HSM imzası üretir"],
        ["GET /signing-sessions/{id}/artifact", "Tamamlanmış imza çıktısını getirir"],
        ["POST /validations/signatures", "CAdES/XAdES/PAdES doğrular"],
        ["POST /validations/certificates", "X.509 sertifika ve zincirini doğrular"],
        ["POST /signatures/cades/augmentations", "CAdES B-LT/B-LTA yükseltir/yeniler"],
        ["/admin/trusted-certificates", "Güven deposu CRUD ve sürümleri"],
        ["/admin/server-key-profiles", "Akıllı kart/HSM profil yönetimi"],
        ["/admin/tsa-profile", "Sürümlü TSA profili"],
        ["/admin/validation-policy", "Sürümlü doğrulama politikası"],
    ], [4200, 5160])

    add_heading(doc, "Ek B. Yardımcı kod parçaları", 1)
    code(doc, r'''
      static void requireSuccess(HttpResponse<String> response) {
          if (response.statusCode() < 200 || response.statusCode() >= 300) {
              throw new IllegalStateException(
                  "E-imza API hatası: HTTP " + response.statusCode() + " - " + response.body());
          }
      }

      static String sha256Hex(byte[] value) throws Exception {
          return HexFormat.of().formatHex(
              MessageDigest.getInstance("SHA-256").digest(value));
      }
    ''', "Java 21")
    code(doc, r'''
      const b64 = bytes => {
        let text = '';
        for (let i = 0; i < bytes.length; i += 32768)
          text += String.fromCharCode(...bytes.slice(i, i + 32768));
        return btoa(text);
      };
      const b64url = bytes => b64(bytes)
        .replaceAll('+','-').replaceAll('/','_').replaceAll('=','');
    ''', "Tarayıcı JavaScript")

    add_heading(doc, "Belge ve proje kaynakları", 1)
    bullets(doc, [
        "OpenAPI: signature-api/src/main/resources/static/openapi/e-signature-api-v1.yaml",
        "Agent OpenAPI: smartcard-agent/src/main/resources/static/openapi/smartcard-agent-v1.yaml",
        "Yönetim/agent rehberi: YONETIM_EKRANI_VE_AGENT_REHBERI.md",
        "Offline masaüstü: desktop-signing-demo/README.md",
        "Standart ve faz kararları: E_IMZA_PROJE_PLANI.md, FAZ_10_UCTAN_UCA_CADES_XADES_PADES.md ve FAZ_11_COKLU_IMZA_VE_JAVA_KUTUPHANE_ENTEGRASYONU.md",
        "Doğrudan JAR örnekleri: JAVA_KUTUPHANE_ENTEGRASYONU.md",
    ])

    doc.core_properties.title = "E-İmza API Entegrasyon Kılavuzu"
    doc.core_properties.subject = "Server-side ve client-side elektronik imza entegrasyonu"
    doc.core_properties.author = "ErbayProject"
    doc.core_properties.keywords = "e-imza, CAdES, XAdES, PAdES, PKCS#11, Smart Card Agent, HSM"
    doc.save(DOCX)
    return DOCX


if __name__ == "__main__":
    print(build())
